"""Insert plaque principle/demo slides after slide 21 (plaque slide) and restyle."""
# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (C) 2026 XMU
#
# This file is part of OralAI-DentalHealth.
# Licensed under the GNU Affero General Public License v3.0 or later.
# See LICENSE and THIRD_PARTY_NOTICES.md in the repository root.
import sys
from pathlib import Path

sys.path.insert(0, f"{_ROOT}/scripts")
from update_ppt_v2 import (OUT_DIR, PPT, add_picture_slide,  # noqa: E402
                           style_run, walk_runs, white_bg)
from pptx import Presentation
from pptx.util import Inches, Pt

# Project root resolved from this file's location (no hardcoded paths).
import os as _os
_ROOT = _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))

FIGS = OUT_DIR / "plaque_figs"


def add_image_grid_slide(prs, title, img_paths, caption=""):
    from PIL import Image
    import cv2
    import numpy as np
    # build one row image
    imgs = [cv2.imread(str(p)) for p in img_paths]
    h = 480
    resized = [cv2.resize(im, (int(im.shape[1] * h / im.shape[0]), h)) for im in imgs]
    sep = 8
    Wc = sum(r.shape[1] for r in resized) + sep * (len(resized) + 1)
    canvas = np.full((h + 2 * sep, Wc, 3), 255, np.uint8)
    x = sep
    for r in resized:
        canvas[sep:sep + h, x:x + r.shape[1]] = r
        x += r.shape[1] + sep
    tmp = OUT_DIR / "plaque_figs/_examples_row.jpg"
    cv2.imwrite(str(tmp), canvas)
    return add_picture_slide(prs, title, tmp, caption)


def main() -> None:
    prs = Presentation(str(PPT))
    n0 = len(prs.slides)
    specs = [
        ("牙菌斑 · 检测原理：牙位级二分类", FIGS / "fig_plaque_principle.jpg",
         "为什么不做像素级分割：标签语义（区域级 0/1）决定任务形式；临床口径即“该牙有无菌斑”"),
        ("牙菌斑 · 模型结构（YOLO11s-cls）", FIGS / "fig_plaque_model_structure.jpg",
         "5.44M 参数 / 12.1 GFLOPs / 输入 224×224 牙位 crop；thr 0.92 在 val 上按筛查优先冻结"),
        ("牙菌斑 · 阈值-指标曲线（extended val 实测）", FIGS / "fig_plaque_threshold.jpg",
         "Sens≥0.95 约束下最大 Spec → 冻结 thr=0.92；test Sens 0.940 / Spec 0.554 / AUROC 0.895"),
        ("牙菌斑 · 模型视角：最高/最低置信度 crop 对比", FIGS / "fig_plaque_crops.jpg",
         "高置信 = 龈缘粉黄色膜面；低置信 = 洁白釉质面——模型学到的判别特征与临床一致"),
        ("牙菌斑 · 更多整图检测示例（实跑）", None,
         "绿框 p<0.92（正常）/ 红框 p≥0.92（菌斑）；来源：完整 pipeline 实跑输出"),
    ]
    # last spec uses the 3-example row
    ex_imgs = [FIGS / f"fig_plaque_example_{k}.jpg" for k in range(3)]
    for title, img, cap in specs:
        if img is None:
            add_image_grid_slide(prs, title, ex_imgs, cap)
        else:
            add_picture_slide(prs, title, img, cap)

    # move the 5 new slides (at end) to right after slide 21 (index 20)
    xml_slides = prs.slides._sldIdLst
    elems = list(xml_slides)
    new_elems = elems[n0:]
    for k, el in enumerate(new_elems):
        xml_slides.remove(el)
        xml_slides.insert(21 + k, el)

    # restyle everything (fonts/black/white bg) once more
    for slide in prs.slides:
        white_bg(slide)
        walk_runs(slide.shapes, black_all=False)
    prs.save(str(PPT))
    print("saved, slides:", len(prs.slides))


if __name__ == "__main__":
    main()

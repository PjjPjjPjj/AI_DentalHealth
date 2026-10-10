"""Plaque TEST-RESULT table: metric | measured value | interpretation."""
# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (C) 2026 XMU
#
# This file is part of OralAI-DentalHealth.
# Licensed under the GNU Affero General Public License v3.0 or later.
# See LICENSE and THIRD_PARTY_NOTICES.md in the repository root.
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.font_manager as fm
import matplotlib.pyplot as plt

# Project root resolved from this file's location (no hardcoded paths).
import os as _os
_ROOT = _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))

ROOT = Path(f"{_ROOT}")
OUT = ROOT / "outputs/demo/pipeline_walkthrough/plaque_figs"
PPT = ROOT / "outputs/demo/pipeline_walkthrough/完整检测汇报.pptx"

ROWS = [
    ("评测集口径", "test 2,472 阳 / 56 阴（clean）", "干净 base 框；患者级隔离 split，官方划分"),
    ("AUROC（test clean）", "0.8950（val 0.8676）", "与阈值无关的排序能力，主指标"),
    ("AP / PR-AUC", "0.9967", "阳性占 98% 时 PR 天然偏高，参考项"),
    ("判定阈值", "P ≥ 0.92（val 上冻结）", "筛查优先：Sens≥0.95 约束下 Spec 最大"),
    ("Sensitivity 召回", "0.9397", "每 100 处菌斑约漏 6 处——核心指标，n=2472 坚实"),
    ("Specificity 特异度", "0.5536", "阴性仅 56 个，±13pt 置信区间，看趋势不抠数值"),
    ("Precision 精确率", "0.9894", "判阳框里 99% 真阳性（类别比例 98:2 所致）"),
    ("F1", "0.9639", "受类别比例影响大，仅供参考"),
    ("extended 口径", "Sens 0.9397 / Spec 0.5429（阴 420）", "纳入增强变体阴性后结论一致，佐证稳健"),
    ("牙位 crop 路径 AUROC", "0.9164（n=282 配对）", "接入牙齿分割的整牙 crop 反而高于 GT 框直推 0.868"),
    ("阴性框概率中位数", "val 0.746 / test 0.880", "阴性多为阳性牙内子区域——天然难的负样本"),
]
HEADERS = ["指标", "实测数值", "解读"]
COL_W = [2.5, 4.0, 6.0]


def draw():
    fm.fontManager.addfont(str(ROOT / "assets/fonts/NotoSansSC.otf"))
    plt.rcParams["font.family"] = "Noto Sans CJK SC"
    n = len(ROWS) + 1
    row_h = 0.52
    fig_h = 0.9 + row_h * n
    fig, ax = plt.subplots(figsize=(14.2, fig_h), dpi=170)
    ax.set_xlim(0, 14.2)
    ax.set_ylim(0, fig_h)
    ax.axis("off")
    xs = [0.25]
    for w in COL_W:
        xs.append(xs[-1] + w)
    y_top = fig_h - 0.55

    def rule(y, lw):
        ax.plot([xs[0], xs[-1]], [y, y], color="black", lw=lw, solid_capstyle="butt")

    rule(y_top + 0.06, 1.8)
    for j, h in enumerate(HEADERS):
        ax.text(xs[j] + 0.12, y_top - row_h / 2, h, fontsize=13.5, weight="bold", va="center")
    rule(y_top - row_h, 0.9)
    for i, row in enumerate(ROWS):
        y = y_top - row_h * (i + 2)
        for j, cell in enumerate(row):
            ax.text(xs[j] + 0.12, y + row_h / 2, cell, fontsize=11.5, va="center", color="#1a1a1a")
        if i < len(ROWS) - 1:
            ax.plot([xs[0], xs[-1]], [y, y], color="#D9D9D9", lw=0.5)
    rule(y_top - row_h * n, 1.8)
    fig.savefig(OUT / "fig_plaque_test_results.jpg", bbox_inches="tight", facecolor="white")
    plt.close(fig)
    print("saved:", OUT / "fig_plaque_test_results.jpg")


def insert_ppt():
    import sys
    sys.path.insert(0, str(ROOT / "scripts"))
    from update_ppt_v2 import add_picture_slide, walk_runs, white_bg  # noqa: E402
    from pptx import Presentation

    prs = Presentation(str(PPT))
    slide = add_picture_slide(
        prs, "牙菌斑 · 测试结果逐项解读（test 集实测）",
        OUT / "fig_plaque_test_results.jpg",
        caption="来源：outputs/plaque_bin_v1/eval_plaque_bin.json；所有数字为真实评测输出")
    xml = prs.slides._sldIdLst
    elems = list(xml)
    new = elems[-1]
    xml.remove(new)
    xml.insert(24, new)   # after 阈值-指标曲线 slide (index 23)
    for s in prs.slides:
        white_bg(s)
        walk_runs(s.shapes, black_all=False)
    prs.save(str(PPT))
    print("PPT saved, slides:", len(prs.slides))


if __name__ == "__main__":
    draw()
    insert_ppt()
